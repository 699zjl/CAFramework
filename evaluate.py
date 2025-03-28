import os
import json
import numpy as np
from copy import deepcopy
from openpyxl import Workbook
from sentence_transformers import SentenceTransformer
from tqdm import tqdm


# A - Z
letter_idx = list(map(chr, range(ord('A'), ord('Z')+1)))

def fast_hist(label_true, label_pred, n_class=2):
    # calculate confusion matrix
    encoding = n_class * label_true.astype(int) + label_pred.astype(int)
    encoding = encoding.flatten()
    hist = np.bincount(encoding, minlength=n_class ** 2)
    hist = hist.reshape(n_class, n_class)
    return hist


def get_all_gts(gts_preds):
    out = set({})
    for img_name in gts_preds:
        gt_pred = gts_preds[img_name]
        out = out | set(gt_pred[0])
    out = sorted(list(out))
    
    # For questions with only two categories, we only calculate the metrics of one category (see function label_level_ss_f1)
    # We fix the order of Yes or No questions to be ['Yes', 'No'] to make the calculated sensitivity and specificity are consistent with common sense.
    if out == ['No', 'Yes']:
        out = ['Yes', 'No']
    return out


def label_level_ss_f1(gt, pred, mask):
    assert gt.shape == pred.shape == mask.shape
    if gt.shape[1] == 1:   
        # for task L1a, which only has one category for GT, calculate sensitiviy
        assert (gt == 1).all()
        _pred = pred[:, 0][mask[:, 0]==1]
        _gt = gt[:, 0][mask[:, 0]==1]
        senss = _pred[_gt==1].sum() / _gt.sum()
        f1s = None
        specs = None      
    elif gt.shape[1] == 2:  
        # for task L2, L3a and L4a, which have two categories for GT, 
        # we only calculate metrics of one category because the metrics of the two categories are symmetric.
        _pred = pred[:, 0][mask[:, 0]==1]
        _gt = gt[:, 0][mask[:, 0]==1]
        senss = _pred[_gt==1].sum() / _gt.sum()
        specs = ((_pred == 0) & (_gt == 0)).sum() / (_gt == 0).sum()
        f1s = 2 * senss * specs / (senss + specs)
    else:
        f1s = []
        senss = []
        specs = []
        for i in range(gt.shape[1]):
            _pred = pred[:, i][mask[:, i]==1]
            _gt = gt[:, i][mask[:, i]==1]

            sens = _pred[_gt==1].sum() / _gt.sum()
            spec = ((_pred == 0) & (_gt == 0)).sum() / (_gt == 0).sum()
            f1 = 2 * sens * spec / (sens + spec)
            f1s.append(f1)
            senss.append(sens)
            specs.append(spec)
        f1s = np.mean(f1s)
        senss = np.mean(senss)
        specs = np.mean(specs)
            
    return f1s, senss, specs
    

class SentenceSimilarity(object):
    
    def __init__(self):
        
        # You can pre-download the model and pass the local path here
        model_path = 'sentence-transformers/all-MiniLM-L6-v2' 
        model_path = '/Users/weiqijie/Desktop/LLMs/code/all-MiniLM-L6-v2'
        self.model = SentenceTransformer(model_path) 
        self.count_all = 0   # The total number of questions
        self.count_match_faild = 0    # The number of questions that answers and options cannot directly matched
        
    def process(self, answers: str, options: list, multiple_answers: bool, bin_mapping: bool, bin_infos: str):
        """
        determine which option the model has chosen
        
        answers: The model output
        options: The options given to the model
        multiple_answers: Whether it is a multi-answer question
        bin_mapping: For Yes/No question, it is unreasonable to directly compare the similarity of options ('Yes' and 'No) and answers (e.g. 'The image has diseases').
                     The Yes/No options need to be converted into their actual meanings (e.g. 'Yes' -> 'abnormality').
        bin_infos: The information used to converted Yes/No options into their actual meanings. e.g. 'Yes' -> 'diabetic retinopathy'
        
        """
        
        self.count_all += 1
        # process the text of answers
        raw_answers = deepcopy(answers)
        answers = answers.replace('\n', ', ')
        answers = answers.replace('and', ', ')
        answers = answers.strip('()')
        
        if multiple_answers:
            answers = answers.split(',')
        else:
            answers = [answers]
            
        for i in range(len(answers)):
            answers[i] = answers[i].strip()
        
        if multiple_answers:
            out_answers = []
            for answer in answers:
                if answer == '':
                    continue
                if answer in letter_idx:
                    out_answer = answer
                elif answer[0] in letter_idx and answer[1] in ['.', '\n']:  # dirctly match
                    out_answer = answer[0]
                elif answer in options:
                    out_answer = letter_idx[options.index(answer)]
                else:
                    out_answer = self.check_existence(answer, options)
                    if len(out_answer) == 1:
                        out_answer = out_answer[0]
                    else:
                        out_answer = None
                if out_answer is not None:
                    out_answers.append(out_answer)
                    
            if len(out_answers) == 0:
                self.count_match_faild += 1
                embeddings_options = self.model.encode(options)
                embedding_answer = self.model.encode([raw_answers])  
                similarities = self.model.similarity(embeddings_options, embedding_answer)  # shape: len(options) * 1
                similarities = similarities[:, 0]
                out_answer = np.argmax(similarities).item()
                out_answer = letter_idx[out_answer]
                out_answers.append(out_answer)
        else:
            answer = answers[0]
            if answer in letter_idx:
                out_answer = answer
            elif answer[0] in letter_idx and answer[1] in ['.', '\n']:  # dirctly match
                out_answer = answer[0]
            elif answer in options:
                out_answer = letter_idx[options.index(answer)]
            else:
                out_answer = self.check_existence(answer, options)
                if len(out_answer) == 1:
                    out_answer = out_answer[0]
                else:
                    if bin_mapping:
                        modify_options = []
                        for opt in options:
                            assert opt in {'Yes', 'No'}, opt
                            if opt == 'Yes':
                                modify_options.append(bin_infos.capitalize())
                            else:
                                modify_options.append('No {}'.format(bin_infos))
                    else:
                        modify_options = deepcopy(options)
                    embeddings_options = self.model.encode(modify_options)
                    embedding_answer = self.model.encode([answer])  
                    self.count_match_faild += 1
                    similarities = self.model.similarity(embeddings_options, embedding_answer)  # shape: len(options) * 1
                    similarities = similarities[:, 0]
                    out_answer = np.argmax(similarities).item()
                    out_answer = letter_idx[out_answer]
            out_answers = [out_answer]
                        
        return out_answers

    def check_existence(self, target_string, ref_strings):
        # check whether text in ref_strings exists in target_string
        contained = []
        for i in range(len(ref_strings)):
            ref_string = ref_strings[i]
            if ref_string.lower() in target_string.lower():
                contained.append(letter_idx[i])
        
        return contained


def main(gt_root, pred_root): 
    
    out_path = os.path.join(pred_root, 'eval.xlsx')


    # 记录所有评价的数据集，不同任务下不同数据集对应的所有可选项
    datasets = []
    SimilarityProcessor = SentenceSimilarity()

    # load all the tasks
    task_infos = []
    level_task_organization = {}
    for level in os.listdir(gt_root):
        if level.startswith('.') or level == 'README.md':
            continue
        if level not in level_task_organization:
            level_task_organization[level] = {}
        for task in os.listdir(os.path.join(gt_root, level)):
            if task.startswith('.'):
                continue
            if os.path.isdir(os.path.join(gt_root, level, task)):
                if task not in level_task_organization[level]:
                    level_task_organization[level][task] = []
                for subtask in os.listdir(os.path.join(gt_root, level, task)):
                    if subtask.startswith('.'):
                        continue    
                    task_infos.append((level, task, subtask))
                    level_task_organization[level][task].append(subtask)
            else:
                task_infos.append((level, task, None))
                level_task_organization[level][task] = None

    # sort all tasks by level, task and subtask
    task_infos = sorted(task_infos, key=lambda x: int(x[2].split('-')[0][3:]) if x[2] is not None else 0)
    task_infos = sorted(task_infos, key=lambda x: x[1])
    task_infos = sorted(task_infos, key=lambda x: x[0])


    # load gt and pred
    raw_datas = {} 
    for count, task_info in enumerate(task_infos):
        level, task, subtask = task_info
        if subtask is None:
            gt_path = os.path.join(gt_root, level, task)
            pred_path = os.path.join(pred_root, level, task)
        else:
            gt_path = os.path.join(gt_root, level, task, subtask)
            pred_path = os.path.join(pred_root, level, task, subtask)
        
        
        if subtask is None:
            print('{} of {}, processing: {}'.format(count + 1, len(task_infos), task))
        else:
            print('{} of {}, processing: {}'.format(count + 1, len(task_infos), subtask))
            
        with open(gt_path) as fin:
            gt_data = json.load(fin)['data'] 
        
        if not os.path.exists(pred_path):
            continue
        with open(pred_path) as fin:
            pred_data = json.load(fin)
        raw_datas[task_info] = {}
        
        for img_name in tqdm(gt_data):
            gt = gt_data[img_name]['gt']   # the corresponding letters for the correct answers
            gt_raw = gt_data[img_name]['raw_data']['gt']   # the text of the correct answers
            options_raw = gt_data[img_name]['raw_data']['options']
            
            if isinstance(gt_raw, list):
                multi_label = True  # whether it is a multi-label question
                gt = gt.split(',')
                gt_raw = gt_raw
            else:
                multi_label = False
                gt = [gt]
                gt_raw = [gt_raw]
                
            if set(options_raw) == {'Yes', 'No'}:
                # refer to SentenceSimilarity.process for details
                assert task in ['L3a-lesion_recognition', 'L4a-binary_condition_diagnosis.json'], task
                bin_mapping = True
                if task == 'L3a-lesion_recognition':
                    bin_infos = subtask.split('-')[-1][:-5]  # the name of the lesion
                else:
                    bin_infos = 'abnormality'
            else:
                bin_mapping = False
                bin_infos = None
            
            pred = pred_data[img_name]
            pred = SimilarityProcessor.process(pred, options_raw, multiple_answers=multi_label, bin_mapping=bin_mapping, bin_infos=bin_infos)
            pred_raw = []
            for _p in pred:
                pred_raw.append(options_raw[letter_idx.index(_p)])
            #print(gt, gt_raw, pred, pred_raw)
            raw_datas[task_info][img_name] = [gt_raw, pred_raw, options_raw]


    task_metrics = {}
    for task_info in raw_datas:
        task_metrics[task_info] = {}
        gts_preds = raw_datas[task_info]
        all_labels = get_all_gts(gts_preds)     # all categories
        gt_mat = np.zeros((len(gts_preds), len(all_labels)))
        pred_mat = np.zeros((len(gts_preds), len(all_labels)))
        mask = np.zeros((len(gts_preds), len(all_labels)), dtype=int)
        
        for i, img_name in enumerate(gts_preds):
            gt_pred = gts_preds[img_name]
            gt, pred, options = gt_pred
            for _g in gt:
                gt_mat[i, all_labels.index(_g)] = 1
            for _p in pred:
                if _p in all_labels:
                    pred_mat[i, all_labels.index(_p)] = 1
            for _o in options:
                if _o in all_labels:
                    mask[i, all_labels.index(_o)] = 1        
        f1, sens, spec = label_level_ss_f1(gt_mat, pred_mat, mask)
        task_metrics[task_info] = [f1, sens, spec]


    # For emode2, some tasks are not evaluated
    valid_levels = []
    valid_tasks = []
    valid_subtasks = []
    for level in level_task_organization:
        for task in level_task_organization[level]:
            if level_task_organization[level][task] is None:
                if (level, task, None) in task_metrics:
                    valid_levels.append(level)
                    valid_tasks.append(task)
            else:
                for sub_task in level_task_organization[level][task]:
                    if (level, task, sub_task) in task_metrics:
                        valid_levels.append(level)
                        valid_tasks.append(task)
                        valid_subtasks.append(sub_task)
    valid_levels = set(valid_levels)
    valid_tasks = set(valid_tasks)
    valid_subtasks = set(valid_subtasks)


    wb = Workbook()
    ws = wb.active
    ws.title = 'overall'
    ws2 = wb.create_sheet(title='detail')

    # write overall results
    column_count = len(valid_levels) + 2
    mean_results = {'Overall': []}

    for level in sorted(level_task_organization.keys()):
        if level not in valid_levels:
            continue
        ws.cell(row=1, column=column_count).value = level
        start_column = column_count
        if level not in mean_results:
            mean_results[level] = []
        for task in sorted(level_task_organization[level].keys()):
            if task not in valid_tasks:
                continue
            if level_task_organization[level][task] is None:
                ws.cell(row=2, column=column_count).value = task[:-5]   # remove '.json'
                
                if task == 'L1a-coarse_modality_perception.json':
                    idx = 1  # for L1a-coarse_modality_perception, report sensitivity instead of F1
                else:
                    idx = 0
                
                ws.cell(row=3, column=column_count).value = task_metrics[(level, task, None)][idx]
                column_count += 1
                mean_results[level].append(task_metrics[(level, task, None)][idx])
                    
            else:
                ws.cell(row=2, column=column_count).value = task
                F1s = []
                for sub_task in level_task_organization[level][task]:
                    if sub_task not in valid_subtasks:
                        continue
                    
                    F1s.append(task_metrics[(level, task, sub_task)][0])
                ws.cell(row=3, column=column_count).value = np.mean(F1s)                
                column_count += 1
                mean_results[level].append(np.mean(F1s))
        
        ws.merge_cells(start_row=1, end_row=1, start_column=start_column, end_column=column_count - 1)
        
    for level in mean_results:
        if level == 'Overall':
            continue
        mean_results['Overall'].append(np.mean(mean_results[level]))

    ws.cell(row=1, column=1).value = 'Average'
    ws.merge_cells(start_row=1, end_row=1, start_column=1, end_column=len(valid_levels) + 1)
    for i, head in enumerate(mean_results):
        ws.cell(row=2, column=i + 1).value = head
        ws.cell(row=3, column=i + 1).value = np.mean(mean_results[head])


    # write detail results
    metric_names = ['F1', 'Sensitivity', 'Specificity']
    column_count = 1
    for level in sorted(level_task_organization.keys()):
        if level not in valid_levels:
            continue
        ws2.cell(row=1, column=column_count).value = level
        start_column = column_count
        for task in sorted(level_task_organization[level].keys()):
            if task not in valid_tasks:
                continue
            if level_task_organization[level][task] is None:
                ws2.cell(row=2, column=column_count).value = task[:-5]   # remove '.json'
                ws2.merge_cells(start_row=2, end_row=3, start_column=column_count, end_column=column_count + 2)
                
                for i, m in enumerate(metric_names):
                    ws2.cell(row=4, column=column_count + i).value = m
                    ws2.cell(row=5, column=column_count + i).value = task_metrics[(level, task, None)][i]
                column_count += 3
                    
            else:
                ws2.cell(row=2, column=column_count).value = task
                start_column2 = column_count
                for sub_task in sorted(level_task_organization[level][task], key=lambda x: int(x.split('-')[0][3:])):
                    if sub_task not in valid_subtasks:
                        continue
                    ws2.cell(row=3, column=column_count).value = sub_task[:-5]  # remove '.json'
                    ws2.merge_cells(start_row=3, end_row=3, start_column=column_count, end_column=column_count + 2)
                    
                    for i, m in enumerate(metric_names):
                        ws2.cell(row=4, column=column_count + i).value = m
                        ws2.cell(row=5, column=column_count + i).value = task_metrics[(level, task, sub_task)][i]
                    column_count += 3
                ws2.merge_cells(start_row=2, end_row=2, start_column=start_column2, end_column=column_count - 1)
        ws2.merge_cells(start_row=1, end_row=1, start_column=start_column, end_column=column_count-1)

    wb.save(out_path)            
    

if __name__ == '__main__':
    gt_root = 'FunBench'
    pred_root = 'answers/test/Qwen2.5-VL-7B-Instruct/emode3'
    main(gt_root, pred_root)     
